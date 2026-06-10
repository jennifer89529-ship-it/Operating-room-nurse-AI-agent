from pathlib import Path
import csv
import json

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "procedure_instrument_checklist_更正版.xlsx"
OUT_JSON = ROOT / "selected_procedure_summary.json"
OUT_CSV = ROOT / "selected_procedure_items.csv"

TARGETS = [
    "動靜脈血管吻合術（A-V shunt）",
    "人工血管置放術（Port-A）",
    "白內障手術（Cataract）",
]


def main() -> None:
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    summary_sheet = workbook["專案術式總表"]
    item_sheet = workbook["術式器械清單"]

    summary_by_procedure = {}
    for row in summary_sheet.iter_rows(min_row=2, values_only=True):
        if row[0] in TARGETS:
            summary_by_procedure[row[0]] = {
                "procedure": row[0],
                "department": row[1],
                "image": row[2],
                "tray": row[3],
                "item_count": int(row[4]),
                "note": row[5],
                "classes": {},
                "items": [],
            }

    headers = [cell.value for cell in item_sheet[1]]
    rows = []
    for row in item_sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        procedure = record["術式"]
        if procedure not in summary_by_procedure:
            continue
        yolo_class = record["YOLO英文類別"]
        summary_by_procedure[procedure]["classes"][yolo_class] = summary_by_procedure[procedure]["classes"].get(yolo_class, 0) + 1
        summary_by_procedure[procedure]["items"].append(
            {
                "name": record["器械名稱"],
                "quantity": record["數量"],
                "class": yolo_class,
                "source": record["對應器械盤"],
                "source_type": record["來源類型"],
            }
        )
        rows.append(
            [
                record["術式"],
                record["科別"],
                record["圖片檔名"],
                record["對應器械盤"],
                record["器械名稱"],
                record["數量"],
                record["YOLO英文類別"],
                record["YOLO中文類別"],
            ]
        )

    ordered = [summary_by_procedure[target] for target in TARGETS]
    OUT_JSON.write_text(json.dumps(ordered, ensure_ascii=False, indent=2), encoding="utf-8")

    with OUT_CSV.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(["術式", "科別", "圖片檔名", "對應器械盤", "器械名稱", "數量", "YOLO英文類別", "YOLO中文類別"])
        writer.writerows(rows)

    print(json.dumps([{k: v for k, v in item.items() if k != "items"} for item in ordered], ensure_ascii=False, indent=2))
    print(f"items={len(rows)}")


if __name__ == "__main__":
    main()
